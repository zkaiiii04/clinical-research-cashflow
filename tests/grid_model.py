"""Python port of the timeline grid builder in vba/modTimeline.bas.

Excel isn't available where the build is done, so this mirrors the macro's
column plan, line building and formula generation, writes the grid formulas to
an .xlsx, recalculates it with LibreOffice and checks the results. It checks
the design and the formula arithmetic. It is not a substitute for running the
VBA in desktop Excel (see tests/UAT_test_script.md).

Usage:
    python3 -I grid_model.py <v0.8 or v0.9 workbook> <scratch dir>

The input workbook's saved snapshots mix projects (budget/actuals are Esko
Combo, mapping/manpower are 3D Gait; both run Nov 2025 - Mar 2027), so this
uses them as one synthetic project and adds test rows for the store,
charge-to, Rate-not-found and out-of-range cases.
"""
import datetime as dt
import os
import re
import subprocess
import sys

import openpyxl

OFF_IO, OFF_BUDGET, OFF_FIRST = 2, 3, 6  # keep in step with vba/modTimeline.bas
BLOCKS = ["MP", "OM", "EQ", "CO", "OTHER"]


def col_letter(c):
    s = ""
    while c:
        c, m = divmod(c - 1, 26)
        s = chr(65 + m) + s
    return s


def ad(r, c):
    return "%s%d" % (col_letter(c), r)


def norm_key(s):
    s = "".join(" " if ch == "\xa0" else ch for ch in str(s) if ord(ch) >= 32 and ord(ch) != 127)
    return re.sub(r" +", " ", s).strip().upper()


def clean(v):
    return "" if v is None else re.sub(r" +", " ", str(v).replace("\xa0", " ")).strip()


def block_of(code):
    return code if code in ("MP", "OM", "EQ", "CO") else "OTHER"


def io_code_text(io, code):
    return io + " (no code)" if code == "" else io + code


def display_category(cat):
    if cat == "":
        return "Manpower"
    if len(cat) > 3 and re.fullmatch(r"[A-Z]+", cat):
        return cat.capitalize()
    return cat


def fmt_extra(x):
    if abs(x) < 1e-7:
        return ""
    return ("+%r" % x) if x > 0 else ("-%r" % -x)


def add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    return dt.date(d.year + y, m + 1, 1)


class Grid:
    def __init__(self, p_start, p_end, fy_basis=True):
        self.fy = fy_basis
        self.m_start = dt.date(p_start.year, p_start.month, 1)
        end = dt.date(p_end.year, p_end.month, 1)
        self.n_months = (end.year - self.m_start.year) * 12 + end.month - self.m_start.month + 1
        self.lines = []
        self.line_index = {}
        self.block_act = {b: [0.0] * (self.n_months + 1) for b in range(1, 6)}
        self.block_has_act = {b: False for b in range(1, 6)}
        self.p_extra = {b: 0.0 for b in range(1, 6)}
        self.soa = None
        self.build_columns()
        self.y_extra = {(b, y): 0.0 for b in range(1, 6) for y in range(1, self.n_y + 1)}

    def month_date(self, i):
        return add_months(self.m_start, i - 1)

    def month_index(self, d):
        return (d.year - self.m_start.year) * 12 + d.month - self.m_start.month + 1

    def in_range(self, mi):
        return 1 <= mi <= self.n_months

    def period_of(self, d):
        if self.fy:
            yk = d.year - (1 if d.month < 4 else 0)
            qn = ((d.month + 8) % 12) // 3 + 1
            return qn, "FY%s/Q%d" % (str(yk)[-2:], qn), "FY%s Total" % str(yk)[-2:]
        qn = (d.month - 1) // 3 + 1
        return qn, "%d | Q%d" % (d.year, qn), "%d Total" % d.year

    def build_columns(self):
        self.cols = [None]  # 1-based: (kind, month, q, y)
        self.q = [None]     # (label, num, first, last, tot, y)
        self.y = [None]     # [label, col]
        prev = None
        for i in range(1, self.n_months + 1):
            qn, ql, yl = self.period_of(self.month_date(i))
            if ql != prev:
                ys = [k for k in range(1, len(self.y)) if self.y[k][0] == yl]
                if ys:
                    y = ys[0]
                else:
                    self.y.append([yl, None])
                    y = len(self.y) - 1
                self.q.append([ql, qn, len(self.cols), None, None, y])
                prev = ql
            nq = len(self.q) - 1
            self.cols.append(("M", i, nq, None))
            nxt = self.period_of(self.month_date(i + 1))[1] if i < self.n_months else ""
            if nxt != ql:
                self.q[nq][3] = len(self.cols) - 1
                self.cols.append(("Q", None, nq, None))
                self.q[nq][4] = len(self.cols) - 1
        for k in range(1, len(self.y)):
            self.cols.append(("Y", None, None, k))
            self.y[k][1] = len(self.cols) - 1
        self.cols.append(("P", None, None, None))
        self.p_col = len(self.cols) - 1
        self.n_cols = self.p_col
        self.n_q = len(self.q) - 1
        self.n_y = len(self.y) - 1

    def add_line(self, kind, blk, io, code, lk, desc):
        line = dict(kind=kind, block=blk, io=io, code=code, lk=lk, desc=desc, label=desc,
                    iocode="", budget=0.0, has_budget=False, seq=len(self.lines) + 1,
                    vals=[None] * (self.n_months + 1), oor=0.0, row=None)
        self.lines.append(line)
        return line

    # ---------------------------------------------------------------- inputs
    def read_budget(self, rows):
        for io, code, desc, amt in rows:
            io, code, desc = clean(io).upper(), clean(code).upper(), clean(desc)
            if not io:
                continue
            blk = block_of(code)
            key = "|".join([blk, code, io, norm_key(desc)])
            if key not in self.line_index:
                ln = self.add_line("BUDGET_MP" if blk == "MP" else "INPUT", blk, io, code, norm_key(desc), desc)
                ln["iocode"] = io_code_text(io, code)
                ln["has_budget"] = True
                self.line_index[key] = ln
            self.line_index[key]["budget"] += amt or 0
        for ln in self.lines:
            if ln["has_budget"] and abs(ln["budget"]) < 0.005:
                ln["kind"] = "HIDDEN"

    def read_manpower(self, rows, charge_to):
        idx = {}
        for io, sg, month, fte, cat, amt, status in rows:
            io, sg, cat = clean(io), clean(sg), clean(cat).upper()
            if not io:
                continue
            rnf = clean(status).upper() == "RATE NOT FOUND"
            blk = "MP"
            if rnf:
                key = "RNF|%s|%s" % (sg.upper(), io.upper())
            else:
                key = "MPW|%s|%s" % (cat, io.upper())
                blk = charge_to.get((cat, io.upper()), "MP")
            if key not in idx:
                if rnf:
                    ln = self.add_line("RNF", "MP", io.upper(), "MP", "", sg or "Manpower")
                    ln["label"] = (sg or "Manpower") + " (Projected) ⚠ Rate not found"
                else:
                    disp = display_category(cat)
                    ln = self.add_line("MANPOWER", blk, io.upper(), "MP", cat, disp)
                    ln["label"] = disp + " (Projected)" + ("" if blk == "MP" else "  ⚙ manpower → " + blk)
                ln["iocode"] = io
                idx[key] = ln
            ln = idx[key]
            if rnf or month is None:
                continue
            mi = self.month_index(month)
            if self.in_range(mi):
                ln["vals"][mi] = (ln["vals"][mi] or 0) + (amt or 0)
            else:
                oor = [x for x in self.lines if x["kind"] == "MP_OOR" and x["block"] == blk]
                o = oor[0] if oor else self.add_line("MP_OOR", blk, "", "MP", "", "manpower months outside project dates")
                o["label"] = "Unmatched: manpower months outside project dates"
                o["oor"] += amt or 0

    def read_actuals(self, rows):
        for code, month, amt in rows:
            code = clean(code).upper()
            b = BLOCKS.index(block_of(code)) + 1
            self.block_has_act[b] = True
            d = dt.date(month.year, month.month, 1)
            self.soa = d if self.soa is None or d > self.soa else self.soa
            mi = self.month_index(d)
            if self.in_range(mi):
                self.block_act[b][mi] += amt
            else:
                yl = self.period_of(d)[2]
                ys = [k for k in range(1, self.n_y + 1) if self.y[k][0] == yl]
                if ys:
                    self.y_extra[(b, ys[0])] += amt
                else:
                    self.p_extra[b] += amt

    def read_store(self, rows):
        orphans = {}
        for blk, code, io, lk, desc, month, amt in rows:
            key = "|".join([blk, code, io, norm_key(lk)])
            mi = self.month_index(month)
            ln = self.line_index.get(key)
            if ln is not None and ln["kind"] == "INPUT" and self.in_range(mi):
                ln["vals"][mi] = (ln["vals"][mi] or 0) + amt
                continue
            if key not in orphans:
                if ln is None:
                    reason = "no longer in budget"
                elif ln["kind"] == "HIDDEN":
                    reason = "line now nets to $0"
                elif ln["kind"] == "INPUT":
                    reason = "month outside project dates"
                else:
                    reason = "MP lines are not typed"
                u = self.add_line("UNMATCHED", blk, io, code, norm_key(lk), desc)
                u["label"] = "Unmatched: %s (%s)" % (desc, reason)
                u["iocode"] = io_code_text(io, code)
                orphans[key] = u
            u = orphans[key]
            if self.in_range(mi):
                u["vals"][mi] = (u["vals"][mi] or 0) + amt
            else:
                u["oor"] += amt

    def sort_lines(self):
        rank = {"INPUT": 1, "BUDGET_MP": 1, "MANPOWER": 2, "RNF": 3, "UNMATCHED": 8}
        out = []
        for ln in self.lines:
            if ln["kind"] == "HIDDEN":
                continue
            r = rank.get(ln["kind"], 9)
            b = BLOCKS.index(ln["block"]) + 1
            if ln["block"] == "MP" and r < 8:
                ln["sort"] = "%d|0|%s|%d|%06d" % (b, ln["iocode"], r, ln["seq"])
            else:
                ln["sort"] = "%d|%d|%s|%06d" % (b, r, ln["iocode"], ln["seq"])
            out.append(ln)
        self.ord = sorted(out, key=lambda x: x["sort"])

    # ---------------------------------------------------------------- write
    def month_le(self, mi):
        return self.soa is not None and self.month_date(mi) <= self.soa

    def year_le(self, y):
        return any(self.q[q][5] == y and self.month_le(self.cols[self.q[q][2]][1]) for q in range(1, self.n_q + 1))

    def qtot_list(self, r, fc, y):
        return ",".join(ad(r, fc + self.q[q][4] - 1) for q in range(1, self.n_q + 1) if self.q[q][5] == y)

    def y_list(self, r, fc):
        return ",".join(ad(r, fc + self.y[y][1] - 1) for y in range(1, self.n_y + 1))

    def write(self, ws, r0=21, c0=2):
        fc = c0 + OFF_FIRST
        put = lambda r, c, v: ws.cell(row=r, column=c, value=v)
        shown = {b: False for b in range(1, 6)}
        for ln in self.ord:
            shown[BLOCKS.index(ln["block"]) + 1] = True
        for b in range(1, 6):
            shown[b] = shown[b] or self.block_has_act[b]
        r = r0 + 2
        rows = {}
        for b in range(1, 6):
            if not shown[b]:
                continue
            hdr = r
            r += 1
            lf = r
            has_bud = False
            for ln in self.ord:
                if BLOCKS.index(ln["block"]) + 1 == b:
                    ln["row"] = r
                    has_bud = has_bud or ln["has_budget"]
                    r += 1
            rows[b] = dict(hdr=hdr, lf=lf, ll=r - 1, P=r, A=r + 1, V=r + 2, VP=r + 3, has_bud=has_bud)
            r += 4
        rows[6] = dict(P=r, A=r + 1, V=r + 2, VP=r + 3)
        self.rows = rows
        for c in range(1, self.n_cols + 1):
            kind, mi, q, y = self.cols[c]
            sc = fc + c - 1
            if kind == "M":
                put(r0 + 1, sc, self.month_date(mi))
            elif kind == "Q":
                put(r0, sc, self.q[q][0])
                put(r0 + 1, sc, "Q Total")
            elif kind == "Y":
                put(r0 + 1, sc, self.y[y][0])
            else:
                put(r0 + 1, sc, "Project Total")
        for ln in self.ord:
            self.write_line(put, ln, c0, fc)
        for b in range(1, 6):
            if shown[b]:
                put(rows[b]["hdr"], c0, BLOCKS[b - 1])
                self.write_block(put, b, rows[b], c0, fc)
        self.write_total(put, shown, rows, c0, fc)

    def write_line(self, put, ln, c0, fc):
        r = ln["row"]
        put(r, c0, ln["label"])
        put(r, c0 + OFF_IO, ln["iocode"])
        if ln["has_budget"]:
            put(r, c0 + OFF_BUDGET, ln["budget"])
        if ln["kind"] in ("BUDGET_MP", "RNF"):
            return
        if ln["kind"] == "MP_OOR":
            put(r, fc + self.p_col - 1, ln["oor"])
            return
        for c in range(1, self.n_cols + 1):
            kind, mi, q, y = self.cols[c]
            sc = fc + c - 1
            if kind == "M":
                if ln["vals"][mi] is not None:
                    put(r, sc, ln["vals"][mi])
            elif kind == "Q":
                put(r, sc, "=SUM(%s:%s)" % (ad(r, fc + self.q[q][2] - 1), ad(r, fc + self.q[q][3] - 1)))
            elif kind == "Y":
                put(r, sc, "=SUM(%s)" % self.qtot_list(r, fc, y))
            else:
                put(r, sc, "=SUM(%s)%s" % (self.y_list(r, fc), fmt_extra(ln["oor"])))

    def pct(self, ra, rv, sc):
        a, v = ad(ra, sc), ad(rv, sc)
        return '=IF(ROUND(%s-%s,2)=0,"",%s/(%s-%s))' % (a, v, v, a, v)

    def write_block(self, put, b, R, c0, fc):
        if R["has_bud"]:
            put(R["P"], c0 + OFF_BUDGET, "=SUM(%s:%s)" % (ad(R["lf"], c0 + OFF_BUDGET), ad(R["ll"], c0 + OFF_BUDGET)))
        for c in range(1, self.n_cols + 1):
            kind, mi, q, y = self.cols[c]
            sc = fc + c - 1
            wrote = False
            if R["ll"] >= R["lf"]:
                put(R["P"], sc, "=SUM(%s:%s)" % (ad(R["lf"], sc), ad(R["ll"], sc)))
            if kind == "M" and self.month_le(mi):
                put(R["A"], sc, self.block_act[b][mi])
                put(R["V"], sc, "=%s-%s" % (ad(R["A"], sc), ad(R["P"], sc)))
                wrote = True
            elif kind == "Q" and self.month_le(self.cols[self.q[q][2]][1]):
                f, l = fc + self.q[q][2] - 1, fc + self.q[q][3] - 1
                put(R["A"], sc, "=SUM(%s:%s)" % (ad(R["A"], f), ad(R["A"], l)))
                put(R["V"], sc, "=SUM(%s:%s)" % (ad(R["V"], f), ad(R["V"], l)))
                wrote = True
            elif kind == "Y":
                ex = self.y_extra[(b, y)]
                if self.year_le(y) or ex:
                    put(R["A"], sc, "=SUM(%s)%s" % (self.qtot_list(R["A"], fc, y), fmt_extra(ex)))
                    extra = "+(%s-SUM(%s))" % (ad(R["A"], sc), self.qtot_list(R["A"], fc, y)) if ex else ""
                    put(R["V"], sc, "=SUM(%s)%s" % (self.qtot_list(R["V"], fc, y), extra))
                    wrote = True
            elif kind == "P":
                ex = self.p_extra[b]
                put(R["A"], sc, "=SUM(%s)%s" % (self.y_list(R["A"], fc), fmt_extra(ex)))
                if self.soa is not None or ex:
                    extra = "+(%s-SUM(%s))" % (ad(R["A"], sc), self.y_list(R["A"], fc)) if ex else ""
                    put(R["V"], sc, "=SUM(%s)%s" % (self.y_list(R["V"], fc), extra))
                    wrote = True
            if wrote:
                put(R["VP"], sc, self.pct(R["A"], R["V"], sc))

    def write_total(self, put, shown, rows, c0, fc):
        T = rows[6]
        bl = [b for b in range(1, 6) if shown[b]]
        if bl:
            put(T["P"], c0 + OFF_BUDGET, "=SUM(%s)" % ",".join(ad(rows[b]["P"], c0 + OFF_BUDGET) for b in bl))
        for c in range(1, self.n_cols + 1):
            kind, mi, q, y = self.cols[c]
            sc = fc + c - 1
            lp = ",".join(ad(rows[b]["P"], sc) for b in bl)
            la = ",".join(ad(rows[b]["A"], sc) for b in bl)
            if lp:
                put(T["P"], sc, "=SUM(%s)" % lp)
            wrote = False
            if kind == "M" and self.month_le(mi):
                put(T["A"], sc, "=SUM(%s)" % la)
                put(T["V"], sc, "=%s-%s" % (ad(T["A"], sc), ad(T["P"], sc)))
                wrote = True
            elif kind == "Q" and self.month_le(self.cols[self.q[q][2]][1]):
                put(T["A"], sc, "=SUM(%s)" % la)
                put(T["V"], sc, "=SUM(%s:%s)" % (ad(T["V"], fc + self.q[q][2] - 1), ad(T["V"], fc + self.q[q][3] - 1)))
                wrote = True
            elif kind == "Y":
                ex = sum(self.y_extra[(b, y)] for b in range(1, 6))
                if self.year_le(y) or ex:
                    put(T["A"], sc, "=SUM(%s)" % la)
                    extra = "+(%s-SUM(%s))" % (ad(T["A"], sc), self.qtot_list(T["A"], fc, y)) if ex else ""
                    put(T["V"], sc, "=SUM(%s)%s" % (self.qtot_list(T["V"], fc, y), extra))
                    wrote = True
            elif kind == "P":
                ex = sum(self.p_extra.values())
                put(T["A"], sc, "=SUM(%s)" % la)
                if self.soa is not None or ex:
                    extra = "+(%s-SUM(%s))" % (ad(T["A"], sc), self.y_list(T["A"], fc)) if ex else ""
                    put(T["V"], sc, "=SUM(%s)%s" % (self.y_list(T["V"], fc), extra))
                    wrote = True
            if wrote:
                put(T["VP"], sc, self.pct(T["A"], T["V"], sc))


def table_rows(wb, sheet, table, cols):
    ws = wb[sheet]
    t = ws.tables[table]
    cells = ws[t.ref]
    hdr = [c.value for c in cells[0]]
    ix = [hdr.index(c) for c in cols]
    return [tuple(row[i].value for i in ix) for row in cells[1:]]


def as_date(v):
    return dt.date(v.year, v.month, 1) if v is not None else None


def to_date(v):
    if v is None:
        return None
    return v.date() if isinstance(v, dt.datetime) else v


def choose_dates(main_io, map_dates, projects):
    """Same rule as modTimeline.CheckTables/ProjectDates: the project dates come
    from qProjects_SelectedTest (the IO Code Type "Main" row for the Main IO,
    else the first "Main" row); the mapping (SOA budget LOA) dates are only a
    fallback. projects = [(IO Code, IO Code Type, Project Start, Project End)].
    Returns (start, end, note)."""
    for pass_ in (1, 2):
        for io, typ, s, e in projects:
            if str(typ or "").strip().upper() == "MAIN" and (pass_ == 2 or str(io or "").strip().upper() == main_io.upper()):
                if s is not None and e is not None:
                    s, e = to_date(s), to_date(e)
                    note = ""
                    if map_dates[0] is not None and (to_date(map_dates[0]), to_date(map_dates[1])) != (s, e):
                        note = "SOA budget LOA dates differ from the project dates"
                    return s, e, note
    if map_dates[0] is not None and map_dates[1] is not None:
        return to_date(map_dates[0]), to_date(map_dates[1]), "project dates missing; using SOA budget LOA dates"
    raise ValueError("no project dates")


def build_case(src_path, fy_basis):
    wb = openpyxl.load_workbook(src_path, data_only=True)
    mapping = table_rows(wb, "PQ_SelectedProjectMapping", "tblSelectedProjectMapping",
                         ["Main IO", "IO Code", "Project Start", "Project End"])
    main = [m for m in mapping if m[0] == "Yes"][0]
    projects = table_rows(wb, "PQ_Projects_SelectedTest", "qProjects_SelectedTest",
                          ["IO Code", "IO Code Type", "Project Start", "Project End"])
    p_start, p_end, _ = choose_dates(main[1], (main[2], main[3]), projects)
    g = Grid(p_start, p_end, fy_basis)
    g.read_budget(table_rows(wb, "PQ_Budget_SelectedTest", "tblBudgetSelectedTest",
                             ["IO Code", "Expense Category", "Budget Description", "Revised Budget ($)"]))
    mp = [(io, sg, as_date(m), fte, cat, amt, st) for io, sg, m, fte, cat, amt, st in table_rows(
        wb, "PQManpower_SelectedTest", "qManpower_SelectedTest",
        ["IO Code", "Staff Group", "Month", "Monthly FTE", "Rate Category", "Projected Amount ($)", "Rate Status"])]
    # Synthetic test rows: an AHS line charged to OM (with one month after the
    # project end), and a Rate-not-found row.
    mp += [("TF25HIP134MP", "Allied Health / Nurse", dt.date(2026, 6, 1), 0.2, "AHS", 1000.0, "Rate matched"),
           ("TF25HIP134MP", "Allied Health / Nurse", dt.date(2027, 4, 1), 0.2, "AHS", 700.0, "Rate matched"),
           ("TF25HIP134MP", "Allied Health / Nurse", dt.date(2026, 5, 1), 0.5, None, None, "Rate not found")]
    g.read_manpower(mp, {("AHS", "TF25HIP134MP"): "OM"})
    acts = [(c, as_date(m), a) for c, m, a in table_rows(
        wb, "PQ_Actuals_SelectedTest", "tblActualsSelectedTest",
        ["Expense Category", "Posting Month", "Actual Amount ($)"]) if a is not None]
    # Out-of-range actuals: Oct 2025 (the FY25 / 2025 column exists) and Mar 2025 (FY24: no column,
    # so Project Total only on the FY basis)
    acts += [("OM", dt.date(2025, 10, 1), 50.0), ("OM", dt.date(2025, 3, 1), 100.0)]
    g.read_actuals(acts)
    store = [("OM", "OM", "TF25HIP136", "HBRA", "HBRA", dt.date(2026, 1, 1), 500.0),
             ("OM", "OM", "TF25HIP136", "HBRA", "HBRA", dt.date(2026, 2, 1), 0.0),
             ("OM", "OM", "TF25HIP136", "Publications", "Publications", dt.date(2026, 3, 1), 1000.0),
             ("OM", "OM", "TF25HIP136", "HBRA", "HBRA", dt.date(2027, 4, 1), 250.0),
             ("OM", "OM", "TF25HIP136", "therapy  associate service", "Therapy Associate Service",
              dt.date(2026, 9, 1), 2550.0)]
    g.read_store(store)
    g.sort_lines()
    return g, acts


def recalc(g, scratch, name):
    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Main_Dashboard"
    g.write(ws)
    path = os.path.join(scratch, name + ".xlsx")
    out.save(path)
    calc_dir = os.path.join(scratch, "calc")
    os.makedirs(calc_dir, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", calc_dir, path],
                   check=True, capture_output=True, timeout=180)
    return openpyxl.load_workbook(os.path.join(calc_dir, name + ".xlsx"), data_only=True)["Main_Dashboard"]


def check(cond, label, problems):
    print(("ok   " if cond else "FAIL ") + label)
    if not cond:
        problems.append(label)


def tdcs_case(scratch, problems):
    """Regression for incident 2026-10-11 (project end date), using the tDCS
    figures from the user's v0.9 workbook: SOA budget LOA dates 1 Jan 2025 -
    31 Dec 2026 in tblSelectedProjectMapping, IREx project dates 1 Jan 2025 -
    30 Sep 2027 in qProjects_SelectedTest. The timeline must run to Sep 2027
    with no manpower months pushed onto an Unmatched row."""
    io = "TF24HIP134MP"
    p_start, p_end, note = choose_dates("TF24HIP134", (dt.date(2025, 1, 1), dt.date(2026, 12, 31)),
                                        [("TF24HIP134", "Main", dt.datetime(2025, 1, 1), dt.datetime(2027, 9, 30))])
    check((p_start, p_end) == (dt.date(2025, 1, 1), dt.date(2027, 9, 30)),
          "tDCS: timeline uses the project dates (1 Jan 2025 - 30 Sep 2027)", problems)
    check(note != "", "tDCS: SOA LOA vs project date difference is reported", problems)
    fb = choose_dates("TF24HIP134", (dt.date(2025, 1, 1), dt.date(2026, 12, 31)),
                      [("TF24HIP134", "Main", None, None)])
    check(fb[:2] == (dt.date(2025, 1, 1), dt.date(2026, 12, 31)) and fb[2] != "",
          "Fallback: blank project dates -> SOA LOA dates, with a note", problems)

    ahs_rate, ra_rate = 7433.3333, 6916.6667
    mp = []
    m = dt.date(2025, 8, 1)
    for _ in range(15):                                   # AHS (OT/PT): Aug 25 - Oct 26, 0.1 FTE
        mp.append((io, "Allied Health / Nurse", m, 0.1, "AHS (OT/PT)", 0.1 * ahs_rate, "Rate matched"))
        m = dt.date(m.year + (m.month == 12), m.month % 12 + 1, 1)
    ra_fte = [0.1, 0.1, 0.2, 0.1, 0.1, 0.1, 0.2, 0.2, 0.3, 0.3, 0.2, 0.2, 0.1, 0.1, 0.1,
              0.1, 0.1, 0.1, 0.6, 0.6, 0.6, 0.6, 0.6, 0.6]  # RA: Oct 25 - Sep 27
    m = dt.date(2025, 10, 1)
    for f in ra_fte:
        mp.append((io, "Research Assistant", m, f, "RA", f * ra_rate, "Rate matched"))
        m = dt.date(m.year + (m.month == 12), m.month % 12 + 1, 1)

    g = Grid(p_start, p_end, True)
    g.read_budget([("TF24HIP134", "MP", "OT", 42000.0), ("TF24HIP134", "MP", "RA", 10800.0)])
    g.read_manpower(mp, {})
    g.read_actuals([])
    g.read_store([])
    g.sort_lines()
    months = [c for c in g.cols[1:] if c[0] == "M"]
    check(len(months) == 33, "tDCS: 33 month columns (Jan 2025 - Sep 2027)", problems)
    check(not any(ln["kind"] == "MP_OOR" for ln in g.ord), "tDCS: no Unmatched manpower row", problems)
    ra = [ln for ln in g.ord if ln["label"] == "RA (Projected)"]
    ahs = [ln for ln in g.ord if ln["label"] == "AHS (OT/PT) (Projected)"]
    check(len(ra) == 1 and abs(sum(v or 0 for v in ra[0]["vals"]) - 43575) < 0.05,
          "tDCS: RA (Projected) months sum to 43,575", problems)
    check(len(ahs) == 1 and abs(sum(v or 0 for v in ahs[0]["vals"]) - 11150) < 0.05,
          "tDCS: AHS (OT/PT) (Projected) months sum to 11,150", problems)
    ws = recalc(g, scratch, "tdcs")
    pt = 2 + OFF_FIRST + g.p_col - 1
    mp_row = g.rows[BLOCKS.index("MP") + 1]["P"]
    check(abs((ws.cell(row=mp_row, column=pt).value or 0) - 54725) < 0.05,
          "tDCS: MP Projected Project Total = 54,725", problems)
    check([y[0] for y in g.y[1:]] == ["FY24 Total", "FY25 Total", "FY26 Total", "FY27 Total"],
          "tDCS: FY24..FY27 totals", problems)


def main(src, scratch):
    problems = []
    fc = 2 + OFF_FIRST

    g, acts = build_case(src, True)
    ws = recalc(g, scratch, "fy")
    months = [c for c in g.cols[1:] if c[0] == "M"]
    check(len(months) == 17, "FY: 17 month columns", problems)
    check([q[0] for q in g.q[1:]] == ["FY25/Q3", "FY25/Q4", "FY26/Q1", "FY26/Q2", "FY26/Q3", "FY26/Q4"],
          "FY: quarters FY25/Q3 .. FY26/Q4", problems)
    check(g.q[1][3] - g.q[1][2] + 1 == 2, "FY: FY25/Q3 holds 2 months (Nov, Dec)", problems)
    check([y[0] for y in g.y[1:]] == ["FY25 Total", "FY26 Total"], "FY: FY25 Total, FY26 Total", problems)
    hdr = [ws.cell(row=22, column=fc + c - 1).value for c in range(1, g.n_cols + 1)]
    check(hdr[-3:] == ["FY25 Total", "FY26 Total", "Project Total"], "FY: totals at the far right", problems)
    check(hdr[2] == "Q Total" and hdr[6] == "Q Total", "FY: a quarter total after each quarter", problems)

    R = g.rows
    pt = fc + g.p_col - 1
    val = lambda r, c: ws.cell(row=r, column=c).value or 0
    check(abs(val(R[6]["P"], 2 + OFF_BUDGET) - 250000) < 0.01, "Budget PROJECT TOTAL = 250,000.00", problems)
    labels = [ln["label"] for ln in g.ord]
    check(not any("Research Fellow" in s for s in labels), "Zero-net Research Fellow hidden", problems)
    tas = [ln for ln in g.ord if ln["label"] == "Therapy Associate Service"]
    check(len(tas) == 1 and abs(tas[0]["budget"] - 30600) < 0.01, "Therapy Associate Service shown once at 30,600",
          problems)
    mp_manpower = sum(v or 0 for ln in g.ord if ln["kind"] == "MANPOWER" and ln["block"] == "MP"
                      for v in ln["vals"])
    check(abs(mp_manpower - 62250) < 0.01, "RA (Projected) | TF25HIP134MP = 62,250.00", problems)
    check(abs(val(R[1]["P"], pt) - 62250) < 0.01, "MP Projected Project Total = 62,250.00", problems)
    ra = [ln for ln in g.ord if ln["label"] == "RA (Projected)"]
    check(len(ra) == 1 and ra[0]["iocode"] == "TF25HIP134MP", "Manpower line label/IO: RA (Projected), TF25HIP134MP",
          problems)
    check(any(ln["label"].startswith("AHS (Projected)") and ln["block"] == "OM" for ln in g.ord),
          "Charge-to: AHS line shown in OM", problems)
    check(any(ln["kind"] == "RNF" for ln in g.ord), "Rate-not-found row shown", problems)

    act_total = sum(a for _, _, a in acts)
    act_mp = sum(a for c, _, a in acts if c == "MP")
    act_om = sum(a for c, _, a in acts if c == "OM")
    check(abs(val(R[6]["A"], pt) - act_total) < 0.01, "Actual PROJECT TOTAL = sum of all actuals (%.2f)" % act_total,
          problems)
    check(abs(val(R[1]["A"], pt) - act_mp) < 0.01, "MP Actual = %.2f" % act_mp, problems)
    check(abs(val(R[2]["A"], pt) - act_om) < 0.01, "OM Actual incl. out-of-range = %.2f" % act_om, problems)
    fy25 = fc + g.y[1][1] - 1
    om_fy25_in_range = sum(a for c, m, a in acts if c == "OM" and dt.date(2025, 11, 1) <= m <= dt.date(2026, 3, 1))
    check(abs(val(R[2]["A"], fy25) - (om_fy25_in_range + 50)) < 0.01, "Oct 2025 actual counted in FY25 Total", problems)

    # Projections: manpower in range + charged AHS + store rows (Apr 2027 rows are outside the dates)
    proj_total = 62250.0003 + 1000 + 700 + 500 + 0 + 1000 + 250 + 2550
    check(abs(val(R[6]["P"], pt) - proj_total) < 0.01, "Projected PROJECT TOTAL = %.2f" % proj_total, problems)
    unmatched = [ln["label"] for ln in g.ord if ln["kind"] == "UNMATCHED"]
    check("Unmatched: Publications (no longer in budget)" in unmatched, "Renamed line -> Unmatched (no longer in budget)",
          problems)
    check("Unmatched: HBRA (month outside project dates)" in unmatched, "Out-of-range month -> Unmatched", problems)
    hbra = [ln for ln in g.ord if ln["label"] == "HBRA"][0]
    check(hbra["vals"][g.month_index(dt.date(2026, 2, 1))] == 0.0, "Saved 0 loads as 0 (not blank)", problems)
    check(hbra["vals"][g.month_index(dt.date(2026, 4, 1))] is None, "Unsaved month stays blank", problems)
    check(any(ln["label"] == "Therapy Associate Service" and ln["vals"][g.month_index(dt.date(2026, 9, 1))] == 2550
              for ln in g.ord), "Store line key is normalised (spaces/case)", problems)

    # Variance to the Latest SOA Month only (Aug 2026)
    soa = g.soa
    check(soa == dt.date(2026, 8, 1), "Latest SOA Month = Aug 2026", problems)
    sep = fc + [c for c in range(1, g.n_cols + 1) if g.cols[c][0] == "M" and g.month_date(g.cols[c][1]) == dt.date(2026, 9, 1)][0] - 1
    check(ws.cell(row=R[6]["V"], column=sep).value in (None, ""), "No variance after the Latest SOA Month", problems)
    proj_to_soa = sum(v or 0 for ln in g.ord if ln["kind"] in ("INPUT", "MANPOWER", "UNMATCHED")
                      for i, v in enumerate(ln["vals"]) if i and g.month_date(i) <= soa)
    exp_v = act_total - proj_to_soa
    check(abs(val(R[6]["V"], pt) - exp_v) < 0.01, "Project variance = actuals - projections to SOA (%.2f)" % exp_v,
          problems)
    exp_pct = exp_v / proj_to_soa
    check(abs(val(R[6]["VP"], pt) - exp_pct) < 1e-9, "Project variance %% = %.4f" % exp_pct, problems)
    for b in (1, 2):
        aa, vv = val(R[b]["A"], pt), val(R[b]["V"], pt)
        blk = BLOCKS[b - 1]
        p_soa = sum(v or 0 for ln in g.ord if ln["block"] == blk and ln["kind"] in ("INPUT", "MANPOWER", "UNMATCHED")
                    for i, v in enumerate(ln["vals"]) if i and g.month_date(i) <= soa)
        check(abs(vv - (aa - p_soa)) < 0.01, "%s variance uses projections to SOA only" % blk, problems)

    g2, _ = build_case(src, False)
    check([q[0] for q in g2.q[1:]] == ["2025 | Q4", "2026 | Q1", "2026 | Q2", "2026 | Q3", "2026 | Q4", "2027 | Q1"],
          "CY: quarters 2025 | Q4 .. 2027 | Q1", problems)
    check(g2.q[1][3] - g2.q[1][2] + 1 == 2, "CY: 2025 | Q4 holds 2 months", problems)
    check([y[0] for y in g2.y[1:]] == ["2025 Total", "2026 Total", "2027 Total"], "CY: 2025/2026/2027 Total", problems)
    ws2 = recalc(g2, scratch, "cy")
    pt2 = fc + g2.p_col - 1
    check(abs((ws2.cell(row=g2.rows[6]["A"], column=pt2).value or 0) - act_total) < 0.01,
          "CY: Actual PROJECT TOTAL unchanged", problems)

    tdcs_case(scratch, problems)

    print("RESULT:", "PASS" if not problems else "FAIL (%d)" % len(problems))
    return 0 if not problems else 1


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    sys.exit(main(sys.argv[1], sys.argv[2]))
