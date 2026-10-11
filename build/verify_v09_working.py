"""Check a built v0.9_working workbook against the v0.8 baseline.

Usage:
    python3 -I verify_v09_working.py <v0.8.xlsm> <v0.9_working.xlsm>

Checks:
1. The Power Query M code (customXml/item1.xml -> DataMashup -> Formulas/Section1.m)
   is identical, and every Power Query part (DataMashup, connections, query
   tables, query-table rels, PQ table definitions) is byte-identical.
2. Only the parts the builder is meant to touch differ; everything else is
   byte-identical.
3. The workbook opens with openpyxl (keep_vba=True); the new sheets, tables,
   names, KPI formulas and visibility are as the spec requires; the VBA
   project is untouched.
"""
import base64
import hashlib
import io
import re
import sys
import zipfile

import openpyxl

EXPECTED_CHANGED = {
    "[Content_Types].xml", "xl/workbook.xml", "xl/_rels/workbook.xml.rels",
    "xl/sharedStrings.xml", "xl/worksheets/sheet1.xml", "xl/worksheets/sheet14.xml",
}
EXPECTED_REMOVED = {
    "xl/calcChain.xml", "xl/comments1.xml", "xl/drawings/vmlDrawing1.vml",
    "xl/worksheets/_rels/sheet1.xml.rels",
}
EXPECTED_ADDED = {
    "xl/worksheets/sheet18.xml", "xl/worksheets/sheet19.xml", "xl/worksheets/sheet20.xml",
    "xl/worksheets/_rels/sheet18.xml.rels", "xl/worksheets/_rels/sheet19.xml.rels",
    "xl/worksheets/_rels/sheet20.xml.rels", "xl/tables/table18.xml", "xl/tables/table19.xml",
}


def m_code(z):
    raw = z.read("customXml/item1.xml")
    txt = raw.decode("utf-16") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else raw.decode("utf-8")
    data = base64.b64decode(re.search(r"<DataMashup[^>]*>(.*)</DataMashup>", txt, re.S).group(1))
    pkg_len = int.from_bytes(data[4:8], "little")
    inner = zipfile.ZipFile(io.BytesIO(data[8:8 + pkg_len]))
    return inner.read("Formulas/Section1.m").decode("utf-8-sig")


def is_pq_part(name, z):
    if name.startswith(("customXml/", "xl/queryTables/")) or name == "xl/connections.xml":
        return True
    if re.match(r"xl/tables/(_rels/)?table\d+\.xml(\.rels)?$", name):
        return b'tableType="queryTable"' in z.read(name) or name.endswith(".rels")
    return False


def main(old_path, new_path):
    problems = []
    zo, zn = zipfile.ZipFile(old_path), zipfile.ZipFile(new_path)
    old, new = set(zo.namelist()), set(zn.namelist())

    mo, mn = m_code(zo), m_code(zn)
    print("M code sha256 v0.8 :", hashlib.sha256(mo.encode()).hexdigest())
    print("M code sha256 v0.9 :", hashlib.sha256(mn.encode()).hexdigest())
    if mo != mn:
        problems.append("Power Query M code differs")
    queries = re.findall(r'shared\s+(#?"[^"]+"|\w+)\s*=', mn)
    print("Queries (%d): %s" % (len(queries), ", ".join(queries)))

    pq_parts = sorted(n for n in old if is_pq_part(n, zo))
    for n in pq_parts:
        if n not in new or zo.read(n) != zn.read(n):
            problems.append("Power Query part changed or missing: " + n)
    print("Power Query parts byte-identical: %d" % len(pq_parts))

    if old - new != EXPECTED_REMOVED:
        problems.append("Unexpected removed parts: %s" % sorted((old - new) ^ EXPECTED_REMOVED))
    if new - old != EXPECTED_ADDED:
        problems.append("Unexpected added parts: %s" % sorted((new - old) ^ EXPECTED_ADDED))
    changed = {n for n in old & new if zo.read(n) != zn.read(n)}
    if changed - EXPECTED_CHANGED:
        problems.append("Unexpected changed parts: %s" % sorted(changed - EXPECTED_CHANGED))
    print("Changed parts: %s" % sorted(changed))
    if zo.read("xl/vbaProject.bin") != zn.read("xl/vbaProject.bin"):
        problems.append("vbaProject.bin changed")

    wb = openpyxl.load_workbook(new_path, keep_vba=True)
    states = {ws.title: ws.sheet_state for ws in wb.worksheets}
    for name, state in [("Main_Dashboard", "visible"), ("Legacy_Timeline", "hidden"),
                        ("Dashboard_Settings", "visible"), ("ProjectionStore", "veryHidden")]:
        if states.get(name) != state:
            problems.append("%s state is %s, expected %s" % (name, states.get(name), state))

    tables = {t: (ws.title, ws.tables[t].ref) for ws in wb.worksheets for t in ws.tables}
    expect_tables = {
        "tblManualProjection": ("Legacy_Timeline", "B59:G60"),
        "tblPatientReimbursement": ("Legacy_Timeline", "B66:H67"),
        "tblManpowerChargeTo": ("Dashboard_Settings", "A4:E5"),
        "tblProjectionStore": ("ProjectionStore", "A1:K2"),
        "qManpower_SelectedTest": ("PQManpower_SelectedTest", "A1:R13"),
    }
    for t, where in expect_tables.items():
        if tables.get(t) != where:
            problems.append("Table %s at %s, expected %s" % (t, tables.get(t), where))

    names = {n: wb.defined_names[n].attr_text for n in wb.defined_names}
    if names.get("TL_Anchor") != "Main_Dashboard!$B$21":
        problems.append("TL_Anchor is %r" % names.get("TL_Anchor"))
    for n in ("TL_TotalBudget", "TL_TotalActual", "TL_TotalProjected"):
        if names.get(n) != '""':
            problems.append("%s is %r" % (n, names.get(n)))

    d = wb["Main_Dashboard"]
    expect = {"C14": "=TL_TotalBudget", "F14": "=TL_TotalActual", "F17": "=TL_TotalProjected",
              "C17": '=IFERROR(F14/C14,"")', "C19": '=IFERROR(F17/C14,"")'}
    for ref, f in expect.items():
        if d[ref].value != f:
            problems.append("Main_Dashboard!%s = %r" % (ref, d[ref].value))
    for ref in ("C7", "C8"):
        if "tblSelectedProjectMapping" not in str(d[ref].value):
            problems.append("Main_Dashboard!%s not repointed" % ref)
    if "tblActualsSelectedTest[Posting Month]" not in str(d["C9"].value):
        problems.append("Main_Dashboard!C9 not repointed")
    for ref in ("B5", "C5", "L14"):
        if d[ref].value is not None:
            problems.append("Main_Dashboard!%s should be empty" % ref)
    if d.max_row > 20:
        problems.append("Main_Dashboard has content below row 20 (max_row %d)" % d.max_row)
    dvs = {str(v.sqref): v.formula1 for v in d.data_validations.dataValidation}
    if "C5" in dvs or dvs.get("C6") != '"Financial Year,Calendar Year"':
        problems.append("Main_Dashboard validations: %s" % dvs)
    for ref in ("E6", "F6", "G6"):
        if d[ref].value is None:
            problems.append("Selector %s missing" % ref)

    leg = wb["Legacy_Timeline"]
    if not str(leg["H25"].value).startswith("=IF(INDEX($21:$21"):
        problems.append("Legacy_Timeline timeline formulas missing")
    if leg["C5"].value is None:
        problems.append("Legacy_Timeline lost its C5 driver")
    if wb["PQManpower_SelectedTest"]["P15"].value is not None:
        problems.append("Stray SUM still in PQManpower_SelectedTest!P15")

    for p in problems:
        print("PROBLEM:", p)
    print("RESULT:", "PASS" if not problems else "FAIL (%d)" % len(problems))
    return 0 if not problems else 1


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    sys.exit(main(sys.argv[1], sys.argv[2]))
